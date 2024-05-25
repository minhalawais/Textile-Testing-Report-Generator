import docx
import os

def read_first_table(docx_file):
    try:
        doc = docx.Document(docx_file)
        table = doc.tables[0]  # Access the first table in the document
        return table
    except Exception as e:
        print(f"Error occurred while reading the first table: {e}")
        return None


def get_files_containing_string(folder_path, part_of_string):
    matching_files = []

    try:
        for filename in os.listdir(folder_path):
            if part_of_string in filename:
                matching_files.append(os.path.join(folder_path, filename))
    except OSError as e:
        print(f"Error occurred while accessing folder '{folder_path}': {e}")

    return matching_files
if __name__ == "__main__":
    folder_path = r"Collage"  # Replace this with the actual folder path
    part_of_string = "05865-23"  # Replace this with the desired part of the string

    matching_files = get_files_containing_string(folder_path, part_of_string)

    if matching_files:
        print("Matching files:")
        for file_path in matching_files:
            if ".docx" in file_path:
                file_path1 = file_path
                print(file_path1)
    else:
        print("No files found matching the specified string.")
    first_table = read_first_table(file_path1)

    if first_table:
        # Print the content of the first table (assuming a basic table structure)
        for i,row in enumerate(first_table.rows):
            for cell in row.cells:
                if cell.text == "P.O:":
                    po_number = first_table.cell(i,1).text
                    print(po_number)
                if cell.text == "Article No:":
                    article_no = first_table.cell(i,1).text
                    print(article_no)
                if cell.text == "Supplier No:":
                    supplier_no = first_table.cell(i,1).text
                    print(supplier_no)
                if cell.text == "Buying Dept (EKB) :":
                    buying_debt = first_table.cell(i,1).text
                    print(buying_debt)
                if "Test Package:" in cell.text:
                    test_package = first_table.cell(i,1).text
                    print(test_package)
    else:
        print("No table found in the document.")
